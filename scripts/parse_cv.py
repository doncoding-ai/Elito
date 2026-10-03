#!/usr/bin/env python3
"""
parse_cv.py — Portfolio auto-updater for Elijah Ndeto.

Drop the latest CV (PDF or DOCX) into cv/. This script:
  1. Finds the newest CV file in cv/
  2. Parses it into structured data
  3. Merges into data/profile.json (only sections it can confidently parse;
     hand-tuned fields like taglines, tags and proficiency levels are preserved)
  4. Copies the newest PDF to cv/latest.pdf for the site's download button

Designed for CVs following Elijah's section structure:
  PROFESSIONAL SUMMARY / CORE COMPETENCIES / TECHNOLOGIES /
  PROFESSIONAL EXPERIENCE / (SELECTED) PROJECTS / EDUCATION

Run:  python scripts/parse_cv.py
Deps: pip install python-docx pdfplumber
"""

import json
import re
import shutil
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CV_DIR = ROOT / "cv"
PROFILE = ROOT / "data" / "profile.json"

SECTION_HEADERS = {
    "PROFESSIONAL SUMMARY": "summary",
    "SUMMARY": "summary",
    "CORE COMPETENCIES": "competencies",
    "TECHNOLOGIES": "technologies",
    "PROFESSIONAL EXPERIENCE": "experience",
    "EXPERIENCE": "experience",
    "SELECTED PROJECTS": "projects",
    "PROJECTS": "projects",
    "EDUCATION": "education",
    "ADDITIONAL INFORMATION": "additional",
}

DATE_RE = re.compile(
    r"((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4}|\d{4})"
    r"\s*[–—-]\s*"
    r"((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4}|\d{4}|Present)",
    re.IGNORECASE,
)
BULLET_RE = re.compile(r"^\s*[•▪▸●·\-–]\s+")


def find_latest_cv() -> Path | None:
    candidates = [
        p for p in CV_DIR.glob("*")
        if p.suffix.lower() in (".pdf", ".docx") and p.name != "latest.pdf"
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda p: p.stat().st_mtime)


def extract_lines(path: Path) -> list[str]:
    if path.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(str(path))
        lines = []
        for para in doc.paragraphs:
            txt = para.text.strip()
            if not txt:
                continue
            # Word list bullets are styling, not text — re-add a marker
            style = (para.style.name if para.style is not None else "").lower()
            is_list = "list" in style or (
                para._p.pPr is not None and para._p.pPr.numPr is not None
            )
            if is_list and not BULLET_RE.match(txt):
                txt = "• " + txt
            lines.append(txt)
        # tables (competencies / technologies live in tables in the docx)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        txt = para.text.strip()
                        if txt:
                            lines.append(txt)
        return lines
    else:
        import pdfplumber
        lines = []
        with pdfplumber.open(str(path)) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                for ln in text.splitlines():
                    ln = ln.strip()
                    if ln:
                        lines.append(ln)
        return lines


def normalise_header(line: str) -> str | None:
    clean = re.sub(r"[^A-Z ]", "", line.upper()).strip()
    return SECTION_HEADERS.get(clean)


def split_sections(lines: list[str]) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {"_head": []}
    current = "_head"
    for ln in lines:
        header = normalise_header(ln)
        if header:
            current = header
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(ln)
    return sections


def parse_experience(lines: list[str]) -> list[dict]:
    """Group lines into roles. A new role starts at a line containing a date range
    (role/company/date may share the line) or at a role-title line followed by
    a company|date line."""
    roles = []
    current = None

    def flush():
        nonlocal current
        if current and current.get("role"):
            roles.append(current)
        current = None

    i = 0
    while i < len(lines):
        ln = lines[i]
        m = DATE_RE.search(ln)
        is_bullet = bool(BULLET_RE.match(ln))

        if m and not is_bullet:
            period = f"{m.group(1)} — {m.group(2)}"
            rest = DATE_RE.sub("", ln).strip(" |·–—-\t")
            role, company = "", ""
            # patterns: "Role · Company", "Company | dates" (role on prev line)
            parts = re.split(r"\s*[·|]\s*", rest)
            parts = [p for p in (p.strip() for p in parts) if p]
            if len(parts) >= 2:
                role, company = parts[0], parts[1]
            elif len(parts) == 1:
                # role was probably the previous non-bullet line
                if current and current.get("role") and not current.get("bullets"):
                    role = current["role"]
                    current = None
                company = parts[0]
                if not role:
                    role = company
                    company = ""
            flush()
            current = {"role": role, "company": company, "period": period,
                       "tag": "", "bullets": [], "stack": []}
        elif is_bullet:
            if current is not None:
                current.setdefault("bullets", []).append(BULLET_RE.sub("", ln).strip())
            # continuation of a wrapped bullet is handled below
        else:
            nxt = lines[i + 1] if i + 1 < len(lines) else ""
            next_is_date = bool(DATE_RE.search(nxt)) and not BULLET_RE.match(nxt)
            if next_is_date:
                # this line is a role title; the next line carries company | dates
                flush()
                current = {"role": ln, "company": "", "period": "",
                           "tag": "", "bullets": [], "stack": []}
            elif current and current.get("bullets"):
                # wrapped bullet continuation
                if len(ln) > 3 and not ln.isupper():
                    current["bullets"][-1] += " " + ln
            else:
                flush()
                current = {"role": ln, "company": "", "period": "",
                           "tag": "", "bullets": [], "stack": []}
        i += 1
    flush()
    return [r for r in roles if r.get("bullets")]


def parse_projects(lines: list[str]) -> list[dict]:
    projects = []
    current = None
    for ln in lines:
        is_continuation = current is not None and (
            BULLET_RE.match(ln) or ln.rstrip().endswith(".") or ln[:1].islower()
        )
        if is_continuation:
            txt = BULLET_RE.sub("", ln).strip()
            current["description"] += (" " if current["description"] else "") + txt
        else:
            if current:
                projects.append(current)
            name = ln
            platform = ""
            m = re.search(r"\(([^)]+)\)", ln)
            if m:
                platform = m.group(1)
                name = re.sub(r"\s*\([^)]+\)", "", ln).strip()
            current = {"name": name, "platform": platform, "description": "", "stack": []}
    if current:
        projects.append(current)
    return [p for p in projects if p["description"]]


def extract_telemetry(all_text: str) -> list[dict]:
    """Pull quantified claims (e.g. '40+ schemas') from the CV text."""
    found = []
    patterns = [
        (r"(\d+)\+?\s*years", "Years in production"),
        (r"(\d+)\+?\s*schemas", "Schemas reverse-engineered"),
        (r"(\d+)\+?\s*source systems", "Source systems architected"),
        (r"migration of (\d+)\+?", "Systems migrated to cloud"),
    ]
    seen = set()
    for pat, label in patterns:
        m = re.search(pat, all_text, re.IGNORECASE)
        if m and label not in seen:
            seen.add(label)
            found.append({"value": int(m.group(1)), "suffix": "+", "label": label, "detail": ""})
    return found


def derive_tag(role: str) -> str:
    r = role.lower()
    if "stream" in r: return "STREAMING"
    if "architect" in r: return "ARCHITECTURE"
    if "model" in r: return "MODELLING"
    if "quality" in r or "auditor" in r: return "QUALITY"
    if "support" in r: return "SUPPORT"
    if "analyst" in r: return "ANALYTICS"
    return "ENGINEERING"


def main() -> int:
    cv = find_latest_cv()
    if not cv:
        print("No CV found in cv/ — nothing to do.")
        return 0

    print(f"Parsing: {cv.name}")
    try:
        lines = extract_lines(cv)
    except Exception as e:
        print(f"ERROR extracting text: {e}", file=sys.stderr)
        return 1

    sections = split_sections(lines)
    profile = json.loads(PROFILE.read_text()) if PROFILE.exists() else {}

    # summary
    if sections.get("summary"):
        profile["summary"] = " ".join(sections["summary"])

    # experience — only overwrite if parse looks sane (>=2 roles with bullets)
    exp = parse_experience(sections.get("experience", []))
    if len(exp) >= 2:
        # preserve hand-tuned tags/stack chips by matching on company+period
        old = {(e.get("company", ""), e.get("period", "")): e
               for e in profile.get("experience", [])}
        old_by_role = {e.get("role", ""): e for e in profile.get("experience", [])}
        for e in exp:
            prev = old.get((e["company"], e["period"])) or old_by_role.get(e["role"])
            if prev:
                e["tag"] = prev.get("tag", "")
                e["stack"] = prev.get("stack", [])
            if not e.get("tag"):
                e["tag"] = derive_tag(e["role"])
        profile["experience"] = exp
        print(f"  experience: {len(exp)} roles")
    else:
        print("  experience: parse inconclusive — keeping existing data")

    # projects
    projects = parse_projects(sections.get("projects", []))
    if projects:
        old = {p.get("name"): p for p in profile.get("projects", [])}
        for p in projects:
            prev = old.get(p["name"])
            if prev:
                p["stack"] = prev.get("stack", [])
        profile["projects"] = projects
        print(f"  projects: {len(projects)}")

    # telemetry
    telemetry = extract_telemetry(" ".join(lines))
    if len(telemetry) >= 3:
        old = {t["label"]: t for t in profile.get("telemetry", [])}
        for t in telemetry:
            if t["label"] in old:
                t["detail"] = old[t["label"]].get("detail", "")
        profile["telemetry"] = telemetry
        print(f"  telemetry: {len(telemetry)} metrics")

    profile["generated_from"] = cv.name
    profile["generated_at"] = date.today().isoformat()

    PROFILE.write_text(json.dumps(profile, indent=2, ensure_ascii=False))
    print(f"Wrote {PROFILE.relative_to(ROOT)}")

    # keep a stable download URL
    if cv.suffix.lower() == ".pdf":
        shutil.copy(cv, CV_DIR / "latest.pdf")
        pub = ROOT / "public" / "cv"
        pub.mkdir(parents=True, exist_ok=True)
        shutil.copy(cv, pub / "latest.pdf")
        print("Updated cv/latest.pdf and public/cv/latest.pdf")
    else:
        print("NOTE: newest CV is a .docx — cv/latest.pdf not updated. "
              "Drop a PDF too if you want the download button current.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
